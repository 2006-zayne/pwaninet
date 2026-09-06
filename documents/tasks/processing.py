"""Celery tasks for document processing pipeline.

These tasks handle asynchronous processing of uploaded documents,
including thumbnail generation, preview creation, metadata extraction,
OCR text extraction, and search indexing. Each stage is independently retryable.

OCR Pipeline Architecture:
- Extension points are provided for future OCR implementation
- OCR text will be stored in DocumentSearchIndex.ocr_text
- OCR text will be indexed with weight D (lower priority)
- The pipeline supports multiple OCR engines (Tesseract, Google Vision, etc.)
"""

import os
import tempfile
from contextlib import contextmanager
from typing import Union
import logging
from celery import shared_task
from celery.exceptions import Retry
from django.conf import settings
from django.utils import timezone

from ..documents.models import Document, DocumentFile, DocumentVersion
from ..services.storage_service import StorageService
from ..services.search_service import SearchService
from ..utils.text_extraction import extract_text_from_file, _clean_and_truncate_text

logger = logging.getLogger(__name__)


@contextmanager
def get_local_filepath(doc_file: DocumentFile):
    """Context manager providing a local filesystem path for a DocumentFile.
    
    If the file is already on local disk, yields its path directly.
    If stored on remote storage (S3/R2), downloads to a temporary file and deletes it on exit.
    """
    if not doc_file.file or not doc_file.file.name:
        yield None
        return

    try:
        path = doc_file.file.path
        if path and os.path.exists(path):
            yield path
            return
    except (NotImplementedError, AttributeError, ValueError):
        pass

    # Check local MEDIA_ROOT fallback (e.g., when remote storage is configured but file exists locally)
    try:
        if hasattr(settings, 'MEDIA_ROOT') and doc_file.file.name:
            local_fallback = os.path.join(str(settings.MEDIA_ROOT), doc_file.file.name)
            if os.path.exists(local_fallback):
                yield local_fallback
                return
    except Exception:
        pass

    # Remote or streamable storage: download to temporary file
    suffix = f".{doc_file.extension}" if doc_file.extension else ""
    temp_path = None
    try:
        try:
            doc_file.file.open('rb')
        except (FileNotFoundError, Exception) as open_err:
            logger.warning(
                "Source file not accessible for file %s (%s): %s",
                doc_file.id, doc_file.file.name, open_err
            )
            yield None
            return

        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tf:
            temp_path = tf.name
            for chunk in doc_file.file.chunks():
                tf.write(chunk)
            doc_file.file.close()
        yield temp_path
    finally:
        if temp_path and os.path.exists(temp_path):
            try:
                os.unlink(temp_path)
            except OSError:
                pass


def extract_ocr_text(file: DocumentFile) -> str:
    """Extract text from document using PyMuPDF, python-docx, python-pptx, etc.
    
    Args:
        file: DocumentFile instance to extract text from
        
    Returns:
        Extracted and sanitized text as string
    """
    try:
        with get_local_filepath(file) as local_path:
            if not local_path:
                return ""
            return extract_text_from_file(
                file_path=local_path,
                mime_type=file.mime_type or "",
                extension=file.extension or ""
            )
    except Exception as exc:
        logger.warning(
            "Failed to extract text from file %s (%s): %s",
            file.id, getattr(file, 'original_filename', ''), exc, exc_info=True
        )
        return ""


@shared_task(queue='docs_queue')
def extract_ocr_text_for_document(document_id_or_instance: Union[int, Document]) -> str:
    """Extract OCR / body text for all files in a document's latest version.
    
    Coordinates extraction across files, enforces guardrails,
    persists extracted text to DocumentSearchIndex.ocr_text,
    and updates weighted search vectors.
    
    Args:
        document_id_or_instance: Document instance or Document ID integer
        
    Returns:
        Combined OCR/body text from all files
    """
    if isinstance(document_id_or_instance, Document):
        document = document_id_or_instance
    else:
        try:
            document = Document.objects.get(id=document_id_or_instance)
        except Document.DoesNotExist:
            logger.warning("Document %s not found for OCR text extraction", document_id_or_instance)
            return ""

    try:
        version = document.latest_version
        if not version:
            return ""
        
        all_text = []
        for file in version.files.all():
            text = extract_ocr_text(file)
            if text:
                all_text.append(text)
        
        combined_text = "\n\n".join(all_text)
        combined_text = _clean_and_truncate_text(combined_text)
        
        # Persist extracted text to DocumentSearchIndex.ocr_text and recompute weighted vectors
        search_service = SearchService()
        search_service.index_document(document, ocr_text=combined_text)
        
        if combined_text:
            logger.info(
                "Extracted %d characters of body text for document %s",
                len(combined_text), document.id
            )
        return combined_text
    except Exception as exc:
        logger.warning(
            "Error extracting OCR text for document %s: %s",
            getattr(document, 'id', document_id_or_instance), exc, exc_info=True
        )
        return ""


@shared_task(bind=True, max_retries=3)
def process_document(self, document_id: int):
    """Main task to process a document after upload.
    
    This task coordinates the entire processing pipeline:
    1. Generate thumbnails
    2. Generate previews
    3. Extract metadata
    4. Extract OCR / body text
    5. Index for search with weighted vectors
    6. Mark as ready
    """
    try:
        document = Document.objects.get(id=document_id)
        logger.info(f"Starting processing for document {document_id}")
        
        # Get latest version
        version = document.latest_version
        if not version:
            logger.error(f"No version found for document {document_id}")
            return
        
        # Process each file in the version
        for file in version.files.all():
            process_file.delay(file.id)
        
        # Wait for all files to be processed (simplified)
        # In production, use group/chord for parallel processing
        
        # Extract OCR / body text and persist to DocumentSearchIndex
        ocr_text = ""
        try:
            ocr_text = extract_ocr_text_for_document(document)
        except Exception as e:
            logger.warning(f"OCR text extraction failed for document {document_id}: {e}", exc_info=True)
        
        # Ensure document is indexed with any updated metadata
        search_service = SearchService()
        search_service.index_document(document, ocr_text=ocr_text if ocr_text else None)
        
        # Mark document as ready
        document.status = 'ready'
        document.save(update_fields=['status'])
        
        logger.info(f"Completed processing for document {document_id}")
        
    except Document.DoesNotExist:
        logger.error(f"Document {document_id} not found")
    except Exception as e:
        logger.error(f"Error processing document {document_id}: {e}")
        document.status = 'draft'
        document.save(update_fields=['status'])
        raise self.retry(exc=e, countdown=60)


@shared_task(bind=True, max_retries=3)
def process_file(self, file_id: int):
    """Process a single document file.
    
    This task handles file-specific processing:
    1. Generate thumbnail
    2. Generate preview
    3. Extract file metadata
    """
    try:
        file = DocumentFile.objects.get(id=file_id)
        logger.info(f"Starting processing for file {file_id}")
        
        file.processing_status = 'processing'
        file.save(update_fields=['processing_status'])
        
        # Generate thumbnail
        generate_thumbnail.delay(file_id)
        
        # Generate preview
        generate_preview.delay(file_id)
        
        # Extract metadata
        extract_metadata.delay(file_id)
        
        # Check for duplicates
        check_duplicate.delay(file_id)
        
        logger.info(f"Completed processing for file {file_id}")
        
    except DocumentFile.DoesNotExist:
        logger.error(f"File {file_id} not found")
    except Exception as e:
        logger.error(f"Error processing file {file_id}: {e}")
        file.processing_status = 'failed'
        file.processing_error = str(e)
        file.save(update_fields=['processing_status', 'processing_error'])
        raise self.retry(exc=e, countdown=60)


@shared_task(bind=True, max_retries=2)
def generate_thumbnail(self, file_id: int):
    """Generate thumbnail for a document file."""
    try:
        file = DocumentFile.objects.get(id=file_id)
        logger.info(f"Generating thumbnail for file {file_id}")
        
        if not file.file:
            logger.info(f"File {file_id} has no file attached, skipping thumbnail")
            return

        # Only generate thumbnails for supported file types
        if file.extension not in ['pdf', 'pptx', 'ppt', 'docx', 'doc']:
            logger.info(f"Skipping thumbnail for {file.extension} file")
            return
        
        from pathlib import Path
        from django.conf import settings
        
        # Create thumbnails directory
        thumbnail_dir = Path(settings.MEDIA_ROOT) / 'thumbnails'
        thumbnail_dir.mkdir(parents=True, exist_ok=True)
        
        # Generate thumbnail based on file type
        if file.extension == 'pdf':
            _generate_pdf_thumbnail(file, thumbnail_dir)
        elif file.extension in ['docx', 'doc']:
            _generate_docx_thumbnail(file, thumbnail_dir)
        elif file.extension in ['pptx', 'ppt']:
            _generate_pptx_thumbnail(file, thumbnail_dir)
        else:
            logger.info(f"Thumbnail generation not supported for {file.extension} file")
            return
        
        logger.info(f"Generated thumbnail for file {file_id}")
        
    except DocumentFile.DoesNotExist:
        logger.error(f"File {file_id} not found")
    except FileNotFoundError as e:
        logger.warning(f"Source file not found for thumbnail of file {file_id}: {e}")
    except Exception as e:
        logger.error(f"Error generating thumbnail for file {file_id}: {e}")
        raise self.retry(exc=e, countdown=30)


@shared_task(bind=True, max_retries=2)
def generate_preview(self, file_id: int):
    """Generate preview for a document file."""
    try:
        file = DocumentFile.objects.get(id=file_id)
        logger.info(f"Generating preview for file {file_id}")
        
        if not file.file:
            logger.info(f"File {file_id} has no file attached, skipping preview")
            file.processing_status = 'failed'
            file.processing_error = 'No file attached'
            file.save(update_fields=['processing_status', 'processing_error'])
            return

        # Generate preview based on file type
        if file.extension == 'pdf':
            _generate_pdf_preview(file)
        elif file.extension in ['docx', 'doc']:
            _generate_docx_preview(file)
        elif file.extension in ['pptx', 'ppt']:
            _generate_pptx_preview(file)
        elif file.extension in ['txt', 'md']:
            _generate_text_preview(file)
        else:
            logger.info(f"Preview generation not supported for {file.extension} file")
            file.processing_status = 'completed'
            file.save(update_fields=['processing_status'])
            return
        
        # Check if preview was generated
        file.refresh_from_db()
        if file.preview_path:
            file.processing_status = 'completed'
            file.processing_error = ''
            file.save(update_fields=['processing_status', 'processing_error'])
            logger.info(f"Generated preview for file {file_id}")
        else:
            file.processing_status = 'failed'
            file.processing_error = file.processing_error or 'Preview generation produced no output'
            file.save(update_fields=['processing_status', 'processing_error'])
            logger.warning(f"Preview was not generated for file {file_id}")
        
    except DocumentFile.DoesNotExist:
        logger.error(f"File {file_id} not found")
    except FileNotFoundError as e:
        logger.warning(f"Source file not found for preview of file {file_id}: {e}")
        try:
            file = DocumentFile.objects.get(id=file_id)
            file.processing_status = 'failed'
            file.processing_error = f"Source file does not exist: {e}"
            file.save(update_fields=['processing_status', 'processing_error'])
        except Exception:
            pass
    except Exception as e:
        logger.error(f"Error generating preview for file {file_id}: {e}")
        # Update status to failed
        try:
            file = DocumentFile.objects.get(id=file_id)
            file.processing_status = 'failed'
            file.processing_error = str(e)
            file.save(update_fields=['processing_status', 'processing_error'])
        except:
            pass
        raise self.retry(exc=e, countdown=30)


def _sync_to_storage_if_remote(local_file_path, rel_storage_path):
    """If default_storage is remote (e.g. S3/R2), upload the locally generated file."""
    try:
        from django.core.files.storage import default_storage
        from django.core.files.base import ContentFile
        if default_storage.__class__.__name__ != 'FileSystemStorage':
            with open(local_file_path, 'rb') as f:
                if default_storage.exists(rel_storage_path):
                    default_storage.delete(rel_storage_path)
                default_storage.save(rel_storage_path, ContentFile(f.read()))
    except Exception as exc:
        logger.warning(f"Failed to sync {rel_storage_path} to remote storage: {exc}")


def _generate_pdf_preview(file):
    """Generate preview for PDF files using PyMuPDF.
    Supports both local and remote storage by using a temporary local copy
    when the storage backend does not expose an absolute path.
    """
    try:
        import fitz  # PyMuPDF
        from pathlib import Path
        from django.conf import settings
        
        preview_dir = Path(settings.MEDIA_ROOT) / 'previews'
        preview_dir.mkdir(parents=True, exist_ok=True)
        
        with get_local_filepath(file) as local_path:
            if not local_path:
                logger.error(f"Unable to obtain local path for PDF preview of file {file.id}")
                return
            doc = fitz.open(local_path)
            if doc.page_count > 0:
                file.page_count = doc.page_count
                page = doc[0]
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                preview_filename = f"{file.id}_preview.jpg"
                preview_path = preview_dir / preview_filename
                pix.save(preview_path)
                doc.close()
                
                rel_path = f"previews/{preview_filename}"
                _sync_to_storage_if_remote(preview_path, rel_path)
                
                file.preview_path = rel_path
                # Also fallback thumbnail_path if missing
                if not file.thumbnail_path:
                    file.thumbnail_path = rel_path
                file.save(update_fields=['preview_path', 'thumbnail_path', 'page_count'])
            else:
                doc.close()
                logger.warning(f"PDF file {file.id} has no pages; preview not generated")
    except Exception as e:
        logger.error(f"Error generating PDF preview for file {file.id}: {e}")


def _generate_docx_preview(file):
    """Generate preview for DOCX files using python-docx."""
    try:
        from docx import Document
        from pathlib import Path
        from django.conf import settings
        from PIL import Image, ImageDraw, ImageFont
        
        preview_dir = Path(settings.MEDIA_ROOT) / 'previews'
        preview_dir.mkdir(parents=True, exist_ok=True)
        
        with get_local_filepath(file) as local_path:
            if not local_path:
                logger.error(f"Unable to obtain local path for DOCX preview of file {file.id}")
                return
            doc = Document(local_path)
            
            img = Image.new('RGB', (800, 600), color='white')
            draw = ImageDraw.Draw(img)
            
            try:
                title_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 24)
            except:
                title_font = ImageFont.load_default()
            
            draw.text((20, 20), "Document Preview", fill='#1F2937', font=title_font)
            
            try:
                text_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 14)
            except:
                text_font = ImageFont.load_default()
            
            y_offset = 60
            for para in doc.paragraphs[:5]:
                text = para.text[:80] + "..." if len(para.text) > 80 else para.text
                draw.text((20, y_offset), text, fill='#374151', font=text_font)
                y_offset += 25
            
            preview_filename = f"{file.id}_preview.jpg"
            preview_path = preview_dir / preview_filename
            img.save(preview_path)
            
            rel_path = f"previews/{preview_filename}"
            _sync_to_storage_if_remote(preview_path, rel_path)
            
            file.preview_path = rel_path
            if not file.thumbnail_path:
                file.thumbnail_path = rel_path
            file.save(update_fields=['preview_path', 'thumbnail_path'])
    except ImportError:
        logger.error("python-docx or PIL not installed")
    except Exception as e:
        logger.error(f"Error generating DOCX preview: {e}")


def _generate_pptx_preview(file):
    """Generate preview for PPTX files using python-pptx."""
    try:
        from pptx import Presentation
        from pathlib import Path
        from django.conf import settings
        from PIL import Image, ImageDraw, ImageFont
        
        preview_dir = Path(settings.MEDIA_ROOT) / 'previews'
        preview_dir.mkdir(parents=True, exist_ok=True)
        
        with get_local_filepath(file) as local_path:
            if not local_path:
                logger.error(f"Unable to obtain local path for PPTX preview of file {file.id}")
                return
            prs = Presentation(local_path)
            img = Image.new('RGB', (800, 600), color='#DC2626')
            draw = ImageDraw.Draw(img)
            
            try:
                title_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 24)
            except:
                title_font = ImageFont.load_default()
            
            draw.text((20, 20), "Presentation Preview", fill='white', font=title_font)
            draw.text((20, 60), f"Slides: {len(prs.slides)}", fill='white', font=title_font)
            
            preview_filename = f"{file.id}_preview.jpg"
            preview_path = preview_dir / preview_filename
            img.save(preview_path)
            
            rel_path = f"previews/{preview_filename}"
            _sync_to_storage_if_remote(preview_path, rel_path)
            
            file.preview_path = rel_path
            if not file.thumbnail_path:
                file.thumbnail_path = rel_path
            file.save(update_fields=['preview_path', 'thumbnail_path'])
    except ImportError:
        logger.error("python-pptx or PIL not installed")
    except Exception as e:
        logger.error(f"Error generating PPTX preview: {e}")


def _generate_text_preview(file):
    """Generate preview for text files."""
    try:
        from pathlib import Path
        from django.conf import settings
        from PIL import Image, ImageDraw, ImageFont
        
        preview_dir = Path(settings.MEDIA_ROOT) / 'previews'
        preview_dir.mkdir(parents=True, exist_ok=True)
        
        with get_local_filepath(file) as local_path:
            if not local_path:
                logger.error(f"Unable to obtain local path for text preview of file {file.id}")
                return
            with open(local_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
            
            img = Image.new('RGB', (800, 600), color='#059669')
            draw = ImageDraw.Draw(img)
            
            try:
                title_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 24)
            except:
                title_font = ImageFont.load_default()
            
            draw.text((20, 20), "Text File Preview", fill='white', font=title_font)
            
            try:
                text_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 14)
            except:
                text_font = ImageFont.load_default()
            
            y_offset = 60
            lines = content.split('\n')[:20]
            for line in lines:
                text = line[:80] + "..." if len(line) > 80 else line
                draw.text((20, y_offset), text, fill='white', font=text_font)
                y_offset += 25
            
            preview_filename = f"{file.id}_preview.jpg"
            preview_path = preview_dir / preview_filename
            img.save(preview_path)
            
            rel_path = f"previews/{preview_filename}"
            _sync_to_storage_if_remote(preview_path, rel_path)
            
            file.preview_path = rel_path
            if not file.thumbnail_path:
                file.thumbnail_path = rel_path
            file.save(update_fields=['preview_path', 'thumbnail_path'])
    except Exception as e:
        logger.error(f"Error generating text preview: {e}")


def _generate_pdf_thumbnail(file, thumbnail_dir):
    """Generate thumbnail for PDF files using PyMuPDF."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        logger.error("PyMuPDF not installed for PDF thumbnail generation")
        if file.preview_path:
            file.thumbnail_path = file.preview_path
            file.save(update_fields=['thumbnail_path'])
        return

    try:
        with get_local_filepath(file) as local_path:
            if not local_path:
                logger.error(f"Unable to obtain local path for PDF thumbnail of file {file.id}")
                return
            doc = fitz.open(local_path)
            if doc.page_count > 0:
                page = doc[0]
                pix = page.get_pixmap(matrix=fitz.Matrix(1, 1))
                thumb_filename = f"{file.id}_thumbnail.jpg"
                thumbnail_path = thumbnail_dir / thumb_filename
                pix.save(thumbnail_path)
                doc.close()
                
                rel_path = f"thumbnails/{thumb_filename}"
                _sync_to_storage_if_remote(thumbnail_path, rel_path)
                
                file.thumbnail_path = rel_path
                file.save(update_fields=['thumbnail_path'])
            else:
                doc.close()
                logger.warning(f"PDF file {file.id} has no pages; thumbnail not generated")
    except Exception as e:
        logger.error(f"Error generating PDF thumbnail: {e}")


def _generate_docx_thumbnail(file, thumbnail_dir):
    """Generate thumbnail for DOCX files."""
    try:
        from PIL import Image, ImageDraw, ImageFont
        
        img = Image.new('RGB', (200, 150), color='#2563EB')
        draw = ImageDraw.Draw(img)
        
        try:
            title_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 20)
        except:
            title_font = ImageFont.load_default()
        
        draw.text((20, 20), "DOCX", fill='white', font=title_font)
        
        thumb_filename = f"{file.id}_thumbnail.jpg"
        thumbnail_path = thumbnail_dir / thumb_filename
        img.save(thumbnail_path)
        
        rel_path = f"thumbnails/{thumb_filename}"
        _sync_to_storage_if_remote(thumbnail_path, rel_path)
        
        file.thumbnail_path = rel_path
        file.save(update_fields=['thumbnail_path'])
    except Exception as e:
        logger.error(f"Error generating DOCX thumbnail: {e}")
        if file.preview_path:
            file.thumbnail_path = file.preview_path
            file.save(update_fields=['thumbnail_path'])


def _generate_pptx_thumbnail(file, thumbnail_dir):
    """Generate thumbnail for PPTX files."""
    try:
        from PIL import Image, ImageDraw, ImageFont
        
        img = Image.new('RGB', (200, 150), color='#DC2626')
        draw = ImageDraw.Draw(img)
        
        try:
            title_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 20)
        except:
            title_font = ImageFont.load_default()
        
        draw.text((20, 20), "PPTX", fill='white', font=title_font)
        
        thumb_filename = f"{file.id}_thumbnail.jpg"
        thumbnail_path = thumbnail_dir / thumb_filename
        img.save(thumbnail_path)
        
        rel_path = f"thumbnails/{thumb_filename}"
        _sync_to_storage_if_remote(thumbnail_path, rel_path)
        
        file.thumbnail_path = rel_path
        file.save(update_fields=['thumbnail_path'])
    except Exception as e:
        logger.error(f"Error generating PPTX thumbnail: {e}")
        if file.preview_path:
            file.thumbnail_path = file.preview_path
            file.save(update_fields=['thumbnail_path'])


@shared_task(bind=True, max_retries=2)
def extract_metadata(self, file_id: int):
    """Extract metadata from a document file."""
    try:
        file = DocumentFile.objects.get(id=file_id)
        logger.info(f"Extracting metadata for file {file_id}")
        
        # This would use libraries like PyPDF2, python-docx, etc.
        # to extract author, title, creation date, etc.
        # For now, this is a placeholder
        
        logger.info(f"Extracted metadata for file {file_id}")
        
    except Exception as e:
        logger.error(f"Error extracting metadata for file {file_id}: {e}")
        raise self.retry(exc=e, countdown=30)


@shared_task(bind=True, max_retries=1)
def check_duplicate(self, file_id: int):
    """Check for duplicate files based on checksum."""
    try:
        file = DocumentFile.objects.get(id=file_id)
        logger.info(f"Checking duplicates for file {file_id}")
        
        from ..storage.models import FileChecksum
        duplicate_files = DocumentFile.objects.filter(
            checksum=file.checksum
        ).exclude(id=file.id)
        
        if duplicate_files.exists():
            logger.warning(f"Found {duplicate_files.count()} duplicates for file {file_id}")
            # Could trigger notification or mark document as potential duplicate
        
        logger.info(f"Checked duplicates for file {file_id}")
        
    except Exception as e:
        logger.error(f"Error checking duplicates for file {file_id}: {e}")
        # Don't retry duplicate check - it's not critical


@shared_task(bind=True, max_retries=1)
def perform_ocr(self, file_id: int):
    """Perform OCR on a document file (future feature).
    
    This task will extract text from images/PDFs for search indexing.
    """
    try:
        file = DocumentFile.objects.get(id=file_id)
        logger.info(f"Performing OCR for file {file_id}")
        
        # This would use libraries like Tesseract, pdf2image, etc.
        # For now, this is a placeholder for future implementation
        
        logger.info(f"Completed OCR for file {file_id}")
        
    except Exception as e:
        logger.error(f"Error performing OCR for file {file_id}: {e}")
        raise self.retry(exc=e, countdown=60)


@shared_task(bind=True, max_retries=1)
def virus_scan(self, file_id: int):
    """Scan a document file for viruses (future feature).
    
    This task will integrate with virus scanning services.
    """
    try:
        file = DocumentFile.objects.get(id=file_id)
        logger.info(f"Scanning file {file_id} for viruses")
        
        # This would integrate with ClamAV or similar
        # For now, this is a placeholder for future implementation
        
        logger.info(f"Completed virus scan for file {file_id}")
        
    except Exception as e:
        logger.error(f"Error scanning file {file_id} for viruses: {e}")
        # Don't retry virus scan - it's not critical


@shared_task
def update_search_index(document_id: int):
    """Update search index for a document."""
    try:
        document = Document.objects.get(id=document_id)
        search_service = SearchService()
        search_service.index_document(document)
        logger.info(f"Updated search index for document {document_id}")
    except Document.DoesNotExist:
        logger.error(f"Document {document_id} not found for search index update")


@shared_task
def remove_from_search_index(document_id: int):
    """Remove a document from the search index."""
    try:
        document = Document.objects.get(id=document_id)
        search_service = SearchService()
        search_service.remove_from_index(document)
        logger.info(f"Removed document {document_id} from search index")
    except Document.DoesNotExist:
        logger.error(f"Document {document_id} not found for search index removal")


@shared_task
def reindex_all_documents():
    """Reindex all documents (maintenance task)."""
    try:
        search_service = SearchService()
        documents = Document.objects.filter(status='ready')
        
        for document in documents:
            search_service.index_document(document)
        
        logger.info(f"Reindexed {documents.count()} documents")
        
    except Exception as e:
        logger.error(f"Error reindexing documents: {e}")


@shared_task
def update_document_analytics(document_id: int):
    """Update cached analytics for a document after engagement.
    
    This task recalculates engagement statistics asynchronously
    to avoid blocking user interactions.
    """
    try:
        from documents.engagement.models import DocumentAnalytics, DocumentView, DocumentDownload, DocumentBookmark, DocumentRating, DocumentShare
        from django.utils import timezone
        from datetime import timedelta
        
        document = Document.objects.get(id=document_id)
        logger.info(f"Updating analytics for document {document_id}")
        
        # Get or create analytics record
        analytics, created = DocumentAnalytics.objects.get_or_create(
            document=document
        )
        
        # Update engagement counts
        analytics.view_count = document.views.count()
        analytics.download_count = document.downloads.count()
        analytics.bookmark_count = document.bookmarks.count()
        analytics.share_count = document.shares.count()
        
        # Update rating statistics
        ratings = document.ratings.all()
        analytics.rating_count = ratings.count()
        analytics.positive_rating_count = ratings.filter(rating=1).count()
        analytics.negative_rating_count = ratings.filter(rating=-1).count()
        
        # Update rating percentages
        analytics.update_rating_percentages()
        
        # Calculate trending score (based on last 7 days)
        seven_days_ago = timezone.now() - timedelta(days=7)
        recent_views = document.views.filter(viewed_at__gte=seven_days_ago).count()
        recent_downloads = document.downloads.filter(downloaded_at__gte=seven_days_ago).count()
        recent_bookmarks = document.bookmarks.filter(created_at__gte=seven_days_ago).count()
        
        # Simple trending algorithm: weighted recent engagement
        analytics.trending_score = (
            recent_views * 1.0 +
            recent_downloads * 2.0 +
            recent_bookmarks * 3.0
        )
        
        # Calculate popularity score (long-term engagement)
        analytics.popularity_score = (
            analytics.view_count * 1.0 +
            analytics.download_count * 2.0 +
            analytics.bookmark_count * 3.0 +
            analytics.positive_rating_count * 5.0
        )
        
        # Update last activity timestamp
        all_activities = []
        if document.views.exists():
            all_activities.append(document.views.latest('viewed_at').viewed_at)
        if document.downloads.exists():
            all_activities.append(document.downloads.latest('downloaded_at').downloaded_at)
        if document.bookmarks.exists():
            all_activities.append(document.bookmarks.latest('created_at').created_at)
        if document.ratings.exists():
            all_activities.append(document.ratings.latest('created_at').created_at)
        if document.shares.exists():
            all_activities.append(document.shares.latest('shared_at').shared_at)
        
        if all_activities:
            analytics.last_activity = max(all_activities)
        
        analytics.save()
        
        logger.info(f"Updated analytics for document {document_id}")
        
    except Document.DoesNotExist:
        logger.error(f"Document {document_id} not found for analytics update")
    except Exception as e:
        logger.error(f"Error updating analytics for document {document_id}: {e}")


@shared_task
def update_analytics_for_all_documents():
    """Update analytics for all documents (maintenance task)."""
    try:
        documents = Document.objects.filter(status='ready')
        
        for document in documents:
            update_document_analytics.delay(document.id)
        
        logger.info(f"Scheduled analytics update for {documents.count()} documents")
        
    except Exception as e:
        logger.error(f"Error scheduling analytics updates: {e}")
