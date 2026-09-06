"""
Document body text extraction pipeline.
Supports PDF (PyMuPDF/fitz), DOCX (python-docx), PPTX (python-pptx),
and plain text/code files with safety caps and sanitization.
"""

import os
import re
import logging
from typing import Optional

logger = logging.getLogger(__name__)

# Safety guardrails
MAX_PAGES = 150
MAX_CHARS = 200_000  # ~35,000 words; prevents memory exhaustion and DB vector limit overflow


def _clean_and_truncate_text(text: str, max_chars: int = MAX_CHARS) -> str:
    """Strip null bytes, normalize repeated whitespace, and apply hard character limit."""
    if not text:
        return ""
    # Strip null bytes (causes PostgreSQL text field errors)
    text = text.replace('\x00', '')
    # Normalize horizontal whitespace
    text = re.sub(r'[ \t]+', ' ', text)
    # Collapse 3+ consecutive newlines to 2
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = text.strip()
    if len(text) > max_chars:
        text = text[:max_chars]
    return text


def _extract_text_from_pdf(file_path: str, max_pages: int = MAX_PAGES) -> str:
    """Extract text from PDF using PyMuPDF (fitz) up to max_pages."""
    import fitz  # PyMuPDF
    text_chunks = []
    doc = None
    try:
        doc = fitz.open(file_path)
        page_limit = min(len(doc), max_pages)
        for page_num in range(page_limit):
            page = doc[page_num]
            page_text = page.get_text("text")
            if page_text:
                text_chunks.append(page_text)
    finally:
        if doc is not None:
            doc.close()
    return "\n".join(text_chunks)


def _extract_text_from_docx(file_path: str) -> str:
    """Extract text from DOCX paragraphs and tables using python-docx."""
    import docx
    doc = docx.Document(file_path)
    text_chunks = []
    
    # Paragraphs
    for p in doc.paragraphs:
        if p.text and p.text.strip():
            text_chunks.append(p.text)
            
    # Tables
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_chunks.append(" | ".join(row_text))
                
    return "\n".join(text_chunks)


def _extract_text_from_pptx(file_path: str) -> str:
    """Extract text from PPTX slides, shapes, and tables using python-pptx."""
    import pptx
    prs = pptx.Presentation(file_path)
    text_chunks = []
    
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    para_text = "".join(run.text for run in paragraph.runs).strip()
                    if para_text:
                        text_chunks.append(para_text)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_chunks.append(" | ".join(row_text))
                        
    return "\n".join(text_chunks)


def _extract_text_from_plain(file_path: str) -> str:
    """Extract text from plain text or source code files."""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1", errors="replace") as f:
            return f.read()


def extract_text_from_file(file_path: str, mime_type: str = "", extension: str = "") -> str:
    """
    Centralized, resilient text extractor for uploaded documents.
    
    Routes by extension or mime_type to the appropriate handler,
    sanitizes content, and enforces guardrails.
    """
    if not file_path or not os.path.isfile(file_path):
        logger.warning("Text extraction skipped: file not found at %s", file_path)
        return ""

    ext = (extension or os.path.splitext(file_path)[1]).lower().lstrip('.')
    mime = (mime_type or '').lower()

    try:
        raw_text = ""
        if ext == 'pdf' or 'pdf' in mime:
            raw_text = _extract_text_from_pdf(file_path)
        elif ext in ('docx', 'doc') or 'wordprocessingml' in mime:
            raw_text = _extract_text_from_docx(file_path)
        elif ext in ('pptx', 'ppt') or 'presentationml' in mime:
            raw_text = _extract_text_from_pptx(file_path)
        elif ext in ('txt', 'csv', 'md', 'py', 'c', 'cpp', 'java', 'js', 'html', 'json', 'xml') or 'text/' in mime:
            raw_text = _extract_text_from_plain(file_path)
        else:
            logger.info("Unsupported extension '%s' (mime: '%s') for text extraction at %s", ext, mime, file_path)
            return ""

        return _clean_and_truncate_text(raw_text)

    except Exception as exc:
        logger.warning("Error extracting text from '%s' (%s): %s", file_path, ext, exc, exc_info=True)
        return ""
