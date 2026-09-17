import os
import tempfile
from unittest.mock import MagicMock, patch
from django.test import SimpleTestCase

import fitz  # PyMuPDF
import docx
from pptx import Presentation

from pwanimate.ingestion.types import (
    ElementType,
    LocationMetadata,
    ExtractedElement,
    ExtractedDocument,
)
from pwanimate.ingestion.exceptions import (
    UnsupportedFormatError,
    CorruptFileError,
    EmptyDocumentError,
)
from pwanimate.ingestion.extractors.factory import (
    extract_document,
    resolve_format,
)
from pwanimate.ingestion.extractors.pdf import PDFExtractor
from pwanimate.ingestion.extractors.docx import DocxExtractor
from pwanimate.ingestion.extractors.pptx import PPTXExtractor
from pwanimate.ingestion.extractors.text import TextExtractor


class PDFExtractorTestCase(SimpleTestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.pdf_path = os.path.join(self.temp_dir.name, "sample.pdf")

        # Synthesize a multi-page PDF using PyMuPDF
        doc = fitz.open()

        # Page 1: Chapter 1 heading + paragraph
        p1 = doc.new_page()
        p1.insert_text((50, 72), "Chapter 1: Operating Systems Overview", fontsize=16)
        p1.insert_text((50, 120), "An operating system acts as an intermediary between the user and hardware.", fontsize=11)

        # Page 2: Section 2 heading + paragraph
        p2 = doc.new_page()
        p2.insert_text((50, 72), "Section 2: Process Scheduling", fontsize=14)
        p2.insert_text((50, 120), "Round-robin scheduling allocates fixed time quanta to processes in the ready queue.", fontsize=11)

        doc.save(self.pdf_path)
        doc.close()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_pdf_multipage_structural_extraction(self):
        extractor = PDFExtractor()
        extracted = extractor.extract(self.pdf_path, filename="sample.pdf")

        self.assertIsInstance(extracted, ExtractedDocument)
        self.assertEqual(extracted.format, "pdf")
        self.assertEqual(extracted.total_pages, 2)
        self.assertGreater(extracted.total_characters, 50)

        # Verify page filtering
        page1_elems = extracted.get_elements_by_page(1)
        page2_elems = extracted.get_elements_by_page(2)

        self.assertGreaterEqual(len(page1_elems), 1)
        self.assertGreaterEqual(len(page2_elems), 1)

        # Page 1 checks
        self.assertEqual(page1_elems[0].location.page_number, 1)
        self.assertIn("Chapter 1", page1_elems[0].text)
        self.assertEqual(page1_elems[0].element_type, ElementType.HEADING)

        # Bounding box presence in extra
        bbox = page1_elems[0].location.extra.get("bbox")
        self.assertIsNotNone(bbox)
        self.assertEqual(len(bbox), 4)

        # Page 2 checks
        self.assertEqual(page2_elems[0].location.page_number, 2)
        self.assertIn("Section 2", page2_elems[0].text)
        self.assertEqual(page2_elems[0].element_type, ElementType.HEADING)

    def test_pdf_empty_document_raises_error(self):
        empty_pdf = os.path.join(self.temp_dir.name, "empty.pdf")
        doc = fitz.open()
        # Page with no text
        doc.new_page()
        doc.save(empty_pdf)
        doc.close()

        extractor = PDFExtractor()
        with self.assertRaises(EmptyDocumentError):
            extractor.extract(empty_pdf)


class DocxExtractorTestCase(SimpleTestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.docx_path = os.path.join(self.temp_dir.name, "lecture.docx")

        # Synthesize a DOCX preserving exact sequence: Heading -> Paragraph -> Table -> Paragraph
        doc = docx.Document()
        doc.add_heading("Database Normalization", level=1)
        doc.add_paragraph("Normalization reduces data redundancy and improves data integrity.")

        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Normal Form"
        table.cell(0, 1).text = "Rule"
        table.cell(1, 0).text = "1NF"
        table.cell(1, 1).text = "Atomic values only"

        doc.add_paragraph("Higher normal forms like 2NF and 3NF address functional dependencies.")
        doc.save(self.docx_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_docx_exact_spatial_sequence_preserved(self):
        extractor = DocxExtractor()
        extracted = extractor.extract(self.docx_path, filename="lecture.docx")

        self.assertIsInstance(extracted, ExtractedDocument)
        self.assertEqual(extracted.format, "docx")
        self.assertEqual(len(extracted.elements), 4)

        # Sequence verification: Heading -> Paragraph -> Table -> Paragraph
        e0, e1, e2, e3 = extracted.elements

        self.assertEqual(e0.element_type, ElementType.HEADING)
        self.assertIn("Database Normalization", e0.text)
        self.assertEqual(e0.location.block_index, 0)

        self.assertEqual(e1.element_type, ElementType.PARAGRAPH)
        self.assertIn("data redundancy", e1.text)
        self.assertEqual(e1.location.section_heading, "Database Normalization")
        self.assertEqual(e1.location.block_index, 1)

        self.assertEqual(e2.element_type, ElementType.TABLE)
        self.assertIn("Normal Form", e2.text)
        self.assertIn("Atomic values only", e2.text)
        self.assertEqual(e2.location.block_index, 2)

        self.assertEqual(e3.element_type, ElementType.PARAGRAPH)
        self.assertIn("Higher normal forms", e3.text)
        self.assertEqual(e3.location.block_index, 3)


class PPTXExtractorTestCase(SimpleTestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.pptx_path = os.path.join(self.temp_dir.name, "lecture.pptx")

        prs = Presentation()

        # Slide 1: Title slide
        title_slide_layout = prs.slide_layouts[0]
        slide1 = prs.slides.add_slide(title_slide_layout)
        slide1.shapes.title.text = "Computer Networks"
        slide1.placeholders[1].text = "Dr. Jane Doe\nDepartment of Computer Science"

        # Slide 2: Bullet slide with speaker note
        bullet_slide_layout = prs.slide_layouts[1]
        slide2 = prs.slides.add_slide(bullet_slide_layout)
        slide2.shapes.title.text = "OSI 7-Layer Model"
        tf = slide2.placeholders[1].text_frame
        tf.text = "Physical Layer"
        p = tf.add_paragraph()
        p.text = "Data Link Layer"

        # Speaker notes for Slide 2
        notes_slide = slide2.notes_slide
        notes_slide.notes_text_frame.text = "IMPORTANT: Focus exam question on the Data Link layer framing."

        prs.save(self.pptx_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_pptx_slides_and_notes_extraction(self):
        extractor = PPTXExtractor()
        extracted = extractor.extract(self.pptx_path, filename="lecture.pptx")

        self.assertIsInstance(extracted, ExtractedDocument)
        self.assertEqual(extracted.format, "pptx")
        self.assertEqual(extracted.total_slides, 2)

        slide1_elems = extracted.get_elements_by_slide(1)
        slide2_elems = extracted.get_elements_by_slide(2)

        self.assertGreaterEqual(len(slide1_elems), 2)
        self.assertGreaterEqual(len(slide2_elems), 2)

        # Slide 1 title check (captured as ElementType.HEADING with is_slide_title metadata)
        self.assertTrue(any(e.element_type == ElementType.HEADING and "Computer Networks" in e.text for e in slide1_elems))

        # Slide 2 speaker note check
        note_elems = [e for e in slide2_elems if e.element_type == ElementType.NOTES]
        self.assertEqual(len(note_elems), 1)
        self.assertIn("exam question", note_elems[0].text)
        self.assertTrue(note_elems[0].metadata.get("is_speaker_notes"))
        self.assertEqual(note_elems[0].location.slide_number, 2)


class TextExtractorTestCase(SimpleTestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.txt_path = os.path.join(self.temp_dir.name, "syllabus.txt")
        content = (
            "# Advanced Algorithms Syllabus\n\n"
            "This course covers graph algorithms, dynamic programming, and complexity.\n\n"
            "```python\n"
            "def dijkstra(graph, start):\n"
            "    pass\n"
            "```\n\n"
            "Final exam will be conducted at the end of the semester.\n"
        )
        with open(self.txt_path, "w", encoding="utf-8") as f:
            f.write(content)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_text_structure_extraction(self):
        extractor = TextExtractor()
        extracted = extractor.extract(self.txt_path, filename="syllabus.txt")

        self.assertIsInstance(extracted, ExtractedDocument)
        self.assertEqual(extracted.format, "txt")
        self.assertGreaterEqual(len(extracted.elements), 3)

        # Heading check
        self.assertEqual(extracted.elements[0].element_type, ElementType.HEADING)
        self.assertIn("Advanced Algorithms Syllabus", extracted.elements[0].text)

        # Code block check
        code_elems = [e for e in extracted.elements if e.element_type == ElementType.CODE]
        self.assertEqual(len(code_elems), 1)
        self.assertIn("def dijkstra", code_elems[0].text)


class FactoryAndSerializationTestCase(SimpleTestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_resolve_format(self):
        self.assertEqual(resolve_format(filename="test.pdf"), "pdf")
        self.assertEqual(resolve_format(filename="test.DOCX"), "docx")
        self.assertEqual(resolve_format(extension="pptx"), "pptx")
        self.assertEqual(resolve_format(mime_type="application/pdf"), "pdf")
        self.assertEqual(resolve_format(mime_type="text/plain; charset=utf-8"), "txt")
        self.assertEqual(resolve_format(mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"), "docx")
        self.assertEqual(resolve_format(filename="unknown.xyz"), "")

    def test_unsupported_format_raises_error(self):
        bad_file = os.path.join(self.temp_dir.name, "program.exe")
        with open(bad_file, "wb") as f:
            f.write(b"\x4d\x5a\x90\x00")

        with self.assertRaises(UnsupportedFormatError):
            extract_document(bad_file)

    def test_corrupt_file_raises_error(self):
        corrupt_pdf = os.path.join(self.temp_dir.name, "bad.pdf")
        with open(corrupt_pdf, "wb") as f:
            f.write(b"NOT A REAL PDF FILE HEADER")

        with self.assertRaises(CorruptFileError):
            extract_document(corrupt_pdf)

    def test_serialization_round_trip(self):
        elem = ExtractedElement(
            text="Distributed Systems Course Notes",
            element_type=ElementType.HEADING,
            location=LocationMetadata(
                page_number=1,
                block_index=0,
                section_heading="Intro",
                extra={"bbox": [10.0, 20.0, 100.0, 50.0]}
            ),
            metadata={"level": 1}
        )
        doc = ExtractedDocument(
            source_filename="notes.pdf",
            format="pdf",
            elements=[elem],
            metadata={"total_pages": 5}
        )

        doc_dict = doc.to_dict()
        restored = ExtractedDocument.from_dict(doc_dict)

        self.assertEqual(restored.source_filename, doc.source_filename)
        self.assertEqual(restored.format, doc.format)
        self.assertEqual(restored.total_pages, doc.total_pages)
        self.assertEqual(len(restored.elements), 1)

        restored_elem = restored.elements[0]
        self.assertEqual(restored_elem.text, elem.text)
        self.assertEqual(restored_elem.element_type, ElementType.HEADING)
        self.assertEqual(restored_elem.location.page_number, 1)
        self.assertEqual(restored_elem.location.extra["bbox"], [10.0, 20.0, 100.0, 50.0])
        self.assertEqual(restored_elem.metadata, {"level": 1})

    def test_extract_document_with_mocked_document_file(self):
        # Create a mock DocumentFile model instance
        mock_file = MagicMock()
        mock_file.id = "doc-123"
        mock_file.original_filename = "notes.txt"
        mock_file.mime_type = "text/plain"
        mock_file.extension = "txt"

        txt_file = os.path.join(self.temp_dir.name, "notes.txt")
        with open(txt_file, "w") as f:
            f.write("Machine Learning Week 1\n\nLinear Regression")

        # Mock documents.tasks.processing.get_local_filepath context manager
        from contextlib import contextmanager

        @contextmanager
        def mock_get_local_filepath(doc_file):
            yield txt_file

        with patch("documents.tasks.processing.get_local_filepath", mock_get_local_filepath):
            extracted = extract_document(mock_file)
            self.assertEqual(extracted.format, "txt")
            self.assertGreaterEqual(len(extracted.elements), 1)
            self.assertIn("Machine Learning Week 1", extracted.elements[0].text)
