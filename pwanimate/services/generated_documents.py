"""Render Pwanimate answers as DOCX/PDF resources in a student's library."""

import hashlib
import base64
import io
import mimetypes
import re
from pathlib import Path

from django.core.files.base import ContentFile
from django.db import transaction
from django.utils.text import slugify

from documents.documents.models import Category, Document, DocumentFile, DocumentVersion
from documents.private_storage import get_private_resource_storage


MAX_RESOURCE_CHARACTERS = 180_000
GENERATED_RENDERER_VERSION = 3
GENERATED_RESOURCE_DESCRIPTION = (
    "A private resource generated from a Pwanimate response. "
    f"Renderer version {GENERATED_RENDERER_VERSION}."
)
_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$")


def _clean_inline(text):
    text = re.sub(r"!?\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*|(?<!_)_([^_]+)_", lambda m: m.group(1) or m.group(2), text)
    return text.strip()


def _split_table_row(line):
    line = line.strip().strip("|")
    cells, cell, escaped = [], [], False
    for char in line:
        if char == "|" and not escaped:
            cells.append("".join(cell).strip())
            cell = []
        else:
            cell.append(char)
        escaped = char == "\\" and not escaped
    cells.append("".join(cell).strip())
    return cells


def _content_blocks(markdown):
    """Parse the Markdown subset used in answers into renderer-neutral blocks."""
    lines, blocks, paragraph = (markdown or "").splitlines(), [], []

    def flush_paragraph():
        if paragraph:
            # Marked is configured with breaks:true in chat, so source newlines
            # within a paragraph must remain explicit line breaks (poems, etc.).
            blocks.append({"type": "paragraph", "text": "\n".join(paragraph)})
            paragraph.clear()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        fence = re.match(r"^\s*(```+|~~~+)(.*)$", line)
        if fence:
            flush_paragraph()
            marker, language = fence.group(1), fence.group(2).strip().split(maxsplit=1)[0] if fence.group(2).strip() else ""
            i += 1
            code = []
            while i < len(lines) and not lines[i].lstrip().startswith(marker):
                code.append(lines[i].rstrip())
                i += 1
            blocks.append({"type": "diagram" if language.lower() in {"mermaid", "flowchart"} else "code", "text": "\n".join(code), "language": language})
            i += 1
            continue
        stripped = line.strip()
        aligned_html = re.fullmatch(r'<p\s+align=["\'](left|right|center)["\']\s*>(.*?)</p>', stripped, re.IGNORECASE)
        if aligned_html:
            flush_paragraph()
            blocks.append({"type": "aligned", "text": aligned_html.group(2), "alignment": aligned_html.group(1).lower()})
            i += 1
            continue
        math_delimiters = (("$$", "$$"), (r"\[", r"\]"))
        matched_math = next(((opening, closing) for opening, closing in math_delimiters if stripped.startswith(opening)), None)
        if matched_math:
            flush_paragraph()
            opening, closing = matched_math
            expression = stripped[len(opening):]
            if expression.endswith(closing) and len(expression) >= len(closing):
                expression = expression[:-len(closing)]
                i += 1
            else:
                math_lines = [expression] if expression else []
                i += 1
                while i < len(lines) and closing not in lines[i]:
                    math_lines.append(lines[i].strip())
                    i += 1
                if i < len(lines):
                    end_line = lines[i].strip()
                    math_lines.append(end_line.split(closing, 1)[0])
                    i += 1
                expression = "\n".join(part for part in math_lines if part)
            if expression.strip():
                blocks.append({"type": "math", "text": expression.strip()})
            continue
        if i + 1 < len(lines) and "|" in line and _TABLE_RULE.match(lines[i + 1]):
            flush_paragraph()
            headers = _split_table_row(line)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip() and "|" in lines[i]:
                cells = _split_table_row(lines[i])
                cells += [""] * max(0, len(headers) - len(cells))
                rows.append(cells[:len(headers)])
                i += 1
            blocks.append({"type": "table", "headers": headers, "rows": rows})
            continue
        if not line.strip():
            flush_paragraph()
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            blocks.append({"type": "heading", "text": heading.group(2), "level": len(heading.group(1))})
        elif re.match(r"^\s*[-*+]\s+", line):
            flush_paragraph()
            blocks.append({"type": "bullet", "text": re.sub(r"^\s*[-*+]\s+", "", line)})
        elif re.match(r"^\s*\d+[.)]\s+", line):
            flush_paragraph()
            blocks.append({"type": "numbered", "text": line.strip()})
        elif line.lstrip().startswith(">"):
            flush_paragraph()
            blocks.append({"type": "quote", "text": line.lstrip()[1:].lstrip()})
        elif re.match(r"^\s*(---+|\*\*\*+)\s*$", line):
            flush_paragraph()
            blocks.append({"type": "rule"})
        else:
            paragraph.append(line.strip())
        i += 1
    flush_paragraph()
    return blocks


def _inline_parts(text):
    """Return safe text/style runs for the common inline Markdown constructs."""
    pattern = re.compile(r"(\\\([^\n]+?\\\)|\\\[[^\n]+?\\\]|(?<!\\)\$[^$\n]+?\$|\*\*.+?\*\*|__.+?__|\*[^*]+\*|_[^_]+_|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))")
    parts, last = [], 0
    for match in pattern.finditer(text or ""):
        if match.start() > last:
            parts.append((text[last:match.start()], ""))
        token = match.group()
        if token.startswith("\\("):
            parts.append((token[2:-2], "math"))
        elif token.startswith("\\["):
            parts.append((token[2:-2], "math"))
        elif token.startswith("$"):
            parts.append((token[1:-1], "math"))
        elif token.startswith("**") or token.startswith("__"):
            parts.append((token[2:-2], "bold"))
        elif token.startswith("*") or token.startswith("_"):
            parts.append((token[1:-1], "italic"))
        elif token.startswith("`"):
            parts.append((token[1:-1], "code"))
        else:
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            parts.append((label, "link:" + url))
        last = match.end()
    if last < len(text or ""):
        parts.append((text[last:], ""))
    return parts or [(text or "", "")]


def _plain_inline(text):
    return "".join(part for part, _style in _inline_parts(text))


def _report_blocks(title, message):
    blocks = _content_blocks(message.content)
    if blocks and blocks[0]["type"] == "heading":
        heading = re.sub(r"\s+", " ", _plain_inline(blocks[0]["text"])).strip().casefold()
        if heading == re.sub(r"\s+", " ", title).strip().casefold():
            blocks = blocks[1:]
    return blocks


def _message_math_expressions(title, message):
    expressions = {}
    for block in _report_blocks(title, message):
        if block["type"] == "math":
            expressions[block["text"]] = True
        for part, style in _inline_parts(block.get("text", "")):
            if style == "math":
                expressions.setdefault(part, False)
    return expressions


def _render_math_images(expressions):
    if not expressions:
        return {}
    from playwright.sync_api import sync_playwright

    katex_js = _static_asset("vendor/katex/katex.min.js").read_text(encoding="utf-8")
    katex_css_path = _static_asset("vendor/katex/katex.min.css")
    katex_css = _inline_css_assets(katex_css_path.read_text(encoding="utf-8"), katex_css_path)
    result = {}
    with sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(channel="chrome", headless=True, args=["--no-sandbox", "--disable-setuid-sandbox"])
        except Exception:
            browser = playwright.chromium.launch(headless=True, args=["--no-sandbox", "--disable-setuid-sandbox"])
        try:
            page = browser.new_page(viewport={"width": 1800, "height": 400}, device_scale_factor=2, color_scheme="light")
            page.set_content(f"<!doctype html><html><head><style>{katex_css}body{{margin:12px;background:white}}.formula{{display:inline-block;padding:4px 8px}}</style></head><body><div id='formulas'></div></body></html>")
            page.add_script_tag(content=katex_js)
            formula_data = [{"expression": expression, "display": display} for expression, display in expressions.items()]
            page.evaluate("formulas => { const root=document.getElementById('formulas'); formulas.forEach(({expression,display},index)=>{ const el=document.createElement('div'); el.className='formula'; el.id='formula-'+index; root.appendChild(el); try { window.katex.render(expression,el,{displayMode:display,throwOnError:false,output:'htmlAndMathml'}); } catch (_) { el.textContent=expression; } }); }", formula_data)
            page.evaluate("async () => { await document.fonts.ready; }")
            for index, (expression, display) in enumerate(expressions.items()):
                image_data = page.locator(f"#formula-{index}").screenshot(type="png", animations="disabled")
                from PIL import Image
                with Image.open(io.BytesIO(image_data)) as image:
                    result[expression] = (image_data, image.width, image.height, display)
        finally:
            browser.close()
    return result


def _flowchart_image(source, scale=1):
    """Draw a safe, intentionally small Mermaid flowchart subset to PNG bytes."""
    from PIL import Image, ImageDraw, ImageFont

    edges, labels, direction = [], {}, "TD"
    for line in (source or "").splitlines():
        line = line.strip()
        if not line or line.startswith("%%"):
            continue
        if line.lower().startswith("flowchart ") or line.lower().startswith("graph "):
            direction = line.split()[-1].upper()
            continue
        match = re.match(r"^([\w-]+)\s*(?:\[([^\]]+)\]|\(([^)]+)\)|\{([^}]+)\})?\s*(?:-->|---|==>)\s*([\w-]+)\s*(?:\[([^\]]+)\]|\(([^)]+)\)|\{([^}]+)\})?$", line)
        if not match:
            continue
        left, la, lb, lc, right, ra, rb, rc = match.groups()
        labels[left] = la or lb or lc or labels.get(left, left)
        labels[right] = ra or rb or rc or labels.get(right, right)
        edges.append((left, right))
    if not labels:
        return None
    nodes = list(labels)
    width, box_w, box_h, gap_x, gap_y = 900, 210, 64, 64, 72
    columns = len(nodes) if direction in {"LR", "RL"} else min(3, len(nodes))
    rows = (len(nodes) + columns - 1) // columns
    height = max(150, 90 + rows * (box_h + gap_y)) if direction not in {"LR", "RL"} else 170
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default()
    positions = {}
    for n, node in enumerate(nodes):
        col, row = (n, 0) if direction in {"LR", "RL"} else (n % columns, n // columns)
        x = 55 + col * (box_w + gap_x)
        y = 42 + row * (box_h + gap_y)
        if x + box_w > width - 30:
            box_w = (width - 70 - gap_x * (columns - 1)) // columns
            x = 35 + col * (box_w + gap_x)
        positions[node] = (x, y)
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=12, fill="#eff6ff", outline="#2563eb", width=2)
        label = labels[node]
        words, lines, current = label.split(), [], ""
        for word in words:
            candidate = (current + " " + word).strip()
            if draw.textlength(candidate, font=font) > box_w - 24 and current:
                lines.append(current); current = word
            else:
                current = candidate
        if current: lines.append(current)
        line_h = 15
        top = y + (box_h - len(lines) * line_h) / 2
        for idx, label_line in enumerate(lines[:3]):
            draw.text((x + box_w / 2, top + idx * line_h), label_line, fill="#172554", font=font, anchor="mt")
    for left, right in edges:
        x1, y1 = positions[left]; x2, y2 = positions[right]
        dx, dy = x2 - x1, y2 - y1
        if abs(dx) > abs(dy):
            start = (x1 + box_w if dx > 0 else x1, y1 + box_h / 2)
            end = (x2 if dx > 0 else x2 + box_w, y2 + box_h / 2)
        else:
            start = (x1 + box_w / 2, y1 + box_h if dy > 0 else y1)
            end = (x2 + box_w / 2, y2 if dy > 0 else y2 + box_h)
        draw.line((start, end), fill="#64748b", width=3)
        if abs(dx) > abs(dy):
            arrow = [(end[0], end[1]), (end[0]-10 if dx > 0 else end[0]+10, end[1]-6), (end[0]-10 if dx > 0 else end[0]+10, end[1]+6)]
        else:
            arrow = [(end[0], end[1]), (end[0]-6, end[1]-10 if dy > 0 else end[1]+10), (end[0]+6, end[1]-10 if dy > 0 else end[1]+10)]
        draw.polygon(arrow, fill="#64748b")
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _static_asset(relative_path):
    from django.contrib.staticfiles import finders

    path = finders.find(relative_path)
    if not path or not isinstance(path, str):
        raise RuntimeError(f"Required static export asset is missing: {relative_path}")
    return Path(path)


def _inline_css_assets(css, css_path):
    """Inline local CSS URLs so the isolated print page needs no web access."""
    url_pattern = re.compile(r"url\((['\"]?)(.*?)\1\)", re.IGNORECASE)

    def replace(match):
        asset = match.group(2).strip()
        if not asset or asset.startswith(("data:", "http:", "https:", "#")):
            return match.group(0)
        asset_path = (css_path.parent / asset.split("?", 1)[0]).resolve()
        if not asset_path.is_file():
            return "none"
        mime = mimetypes.guess_type(asset_path.name)[0] or "application/octet-stream"
        encoded = base64.b64encode(asset_path.read_bytes()).decode("ascii")
        return f'url("data:{mime};base64,{encoded}")'

    return url_pattern.sub(replace, css)


def _prepare_mermaid_markdown(markdown):
    """Replace supported Mermaid flowcharts with embedded diagram images."""
    replacements = {}
    pattern = re.compile(r"^\s*```(?:mermaid|flowchart)\s*\n(.*?)^\s*```\s*$", re.MULTILINE | re.DOTALL | re.IGNORECASE)

    def replace(match):
        image = _flowchart_image(match.group(1))
        if not image:
            return match.group(0)
        token = f"PWANIMATE_DIAGRAM_{len(replacements)}_PLACEHOLDER"
        replacements[token] = "<figure class=\"pwanimate-export-diagram\"><img alt=\"Flowchart\" src=\"data:image/png;base64," + base64.b64encode(image).decode("ascii") + "\"></figure>"
        return f"\n\n{token}\n\n"

    return pattern.sub(replace, markdown or ""), replacements


def _render_markdown_html(title, message):
    """Render answer content with the same Marked, KaTeX, DOMPurify and Prism stack as chat."""
    content = message.content or ""
    assets = {
        "marked": _static_asset("vendor/marked/marked.min.js"),
        "dompurify": _static_asset("vendor/dompurify/purify.min.js"),
        "katex": _static_asset("vendor/katex/katex.min.js"),
        "prism": _static_asset("vendor/prism/prism.min.js"),
        "prism_python": _static_asset("vendor/prism/prism-python.min.js"),
        "prism_sql": _static_asset("vendor/prism/prism-sql.min.js"),
        "prism_bash": _static_asset("vendor/prism/prism-bash.min.js"),
        "prism_json": _static_asset("vendor/prism/prism-json.min.js"),
        "prism_javascript": _static_asset("vendor/prism/components/prism-javascript.min.js"),
        "prism_markup": _static_asset("vendor/prism/components/prism-markup.min.js"),
        "katex_css": _static_asset("vendor/katex/katex.min.css"),
        "prism_css": _static_asset("vendor/prism/prism.min.css"),
        "pwanimate_css": _static_asset("css/pwanimate/pwanimate.css"),
    }
    sources = _source_labels(message)
    if sources:
        content += "\n\n## Sources\n\n" + "\n".join(
            "- " + re.sub(r"([\\`*_{}\[\]()#+.!|>])", r"\\\1", source)
            for source in sources
        )
    markdown, diagram_replacements = _prepare_mermaid_markdown(content)
    first_heading = next((block for block in _content_blocks(content) if block["type"] == "heading"), None)
    if first_heading and re.sub(r"\s+", " ", _plain_inline(first_heading["text"])).strip().casefold() == re.sub(r"\s+", " ", title).strip().casefold():
        markdown = re.sub(r"^#{1,6}\s+.+(?:\n+|$)", "", markdown, count=1)

    # Use static inline assets so export rendering is deterministic and network isolated.
    styles = "\n".join(
        f"<style>{_inline_css_assets(assets[key].read_text(encoding='utf-8'), assets[key])}</style>"
        for key in ("katex_css", "prism_css", "pwanimate_css")
    )
    custom_css = r"""
        :root { color-scheme: light; --primary:#2563eb; --pwanimate-asst-text:#0f172a; --pwanimate-muted:#64748b; --pwanimate-card-bg:#fff; --pwanimate-border:#cbd5e1; --pwanimate-hover-bg:#f1f5f9; }
        html, body { margin:0; padding:0; background:#fff; color:#0f172a; }
        body { font-family: Inter, Aptos, Arial, sans-serif; font-size:10.5pt; line-height:1.58; -webkit-print-color-adjust:exact; print-color-adjust:exact; }
        .export-shell { max-width:72ch; margin:0 auto; }
        .export-title { margin:0 0 1.15em; color:#194a89; font-size:22pt; line-height:1.2; font-weight:700; }
        .pwanimate-markdown-body [align="right"], .pwanimate-markdown-body .text-end { text-align:right; }
        .pwanimate-markdown-body [align="center"], .pwanimate-markdown-body .text-center { text-align:center; }
        .pwanimate-markdown-body .pwanimate-table-scroll { overflow:visible; max-width:none; }
        .pwanimate-markdown-body table { width:100%; min-width:0; max-width:100%; table-layout:auto; border-collapse:collapse; }
        .pwanimate-markdown-body th, .pwanimate-markdown-body td { white-space:normal; overflow-wrap:anywhere; word-break:normal; min-width:0; }
        .pwanimate-markdown-body thead { display:table-header-group; }
        .pwanimate-markdown-body tr { break-inside:avoid; page-break-inside:avoid; }
        .pwanimate-markdown-body .pwanimate-code-block { break-inside:avoid; page-break-inside:avoid; border-radius:9px; box-shadow:none; }
        .pwanimate-markdown-body .pwanimate-code-copy-btn { display:none !important; }
        .pwanimate-markdown-body pre { white-space:pre-wrap; overflow-wrap:anywhere; }
        .pwanimate-markdown-body .katex-display { max-width:100%; overflow:visible; break-inside:avoid; page-break-inside:avoid; }
        .pwanimate-markdown-body .katex-display > .katex { max-width:100%; white-space:normal; }
        .pwanimate-export-diagram { margin:1rem 0; text-align:center; break-inside:avoid; }
        .pwanimate-export-diagram img { max-height:8in; object-fit:contain; margin:0 auto; }
        .wide-table-page { page: wide-table; break-before:page; break-after:page; }
        @page { size:A4 portrait; margin:18mm 19mm 23mm; }
        @page wide-table { size:A4 landscape; margin:15mm 15mm 22mm; }
        @media print {
            .export-shell { max-width:none; }
            .wide-table-page table { font-size:8pt; }
            .wide-table-page.dense-wide-table table { font-size:6.5pt; }
            .wide-table-page th, .wide-table-page td { padding:4pt 5pt; }
        }
    """
    page_html = f"<!doctype html><html><head><meta charset=\"utf-8\"><title></title>{styles}<style>{custom_css}</style></head><body><main class=\"export-shell\"><h1 class=\"export-title\"></h1><article id=\"answer\" class=\"pwanimate-markdown-body\"></article></main></body></html>"
    return page_html, markdown, diagram_replacements


def _source_labels(message):
    labels = []
    for source in message.sources if isinstance(message.sources, list) else []:
        if isinstance(source, str):
            label = source.strip()
        elif isinstance(source, dict):
            label = source.get("title") or source.get("citation") or source.get("url") or ""
        else:
            label = ""
        if label and label not in labels:
            labels.append(label)
    for citation in message.citations if isinstance(message.citations, list) else []:
        label = citation if isinstance(citation, str) else ""
        if label and label not in labels:
            labels.append(label)
    return labels


def _render_docx(title, message):
    from docx import Document as WordDocument
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = WordDocument()
    doc.core_properties.title = title
    doc.core_properties.author = "Pwanimate"
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(40, 48, 60)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(25, 74, 137)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(5)

    math_images = _render_math_images(_message_math_expressions(title, message))
    title_paragraph = doc.add_paragraph(style="Title")
    title_paragraph.add_run(title)
    subtitle = doc.add_paragraph("Prepared with Pwanimate")
    subtitle.style = "Subtitle"

    def shade(cell, fill):
        props = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); props.append(shd)

    def style_run(run, style):
        if style == "bold": run.bold = True
        elif style == "italic": run.italic = True
        elif style == "code":
            run.font.name = "Consolas"; run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(153, 27, 27)
        elif style.startswith("link:"):
            run.underline = True; run.font.color.rgb = RGBColor(37, 99, 235)

    def add_inline(paragraph, text):
        for part, style in _inline_parts(text):
            if style == "math" and part in math_images:
                image_data, image_width, image_height, _display = math_images[part]
                run = paragraph.add_run()
                run.add_picture(io.BytesIO(image_data), width=Inches(min(2.4, image_width / 192)))
            else:
                pieces = part.split("\n")
                for index, piece in enumerate(pieces):
                    if piece:
                        run = paragraph.add_run(piece); style_run(run, style)
                    if index < len(pieces) - 1:
                        paragraph.add_run().add_break()

    for block in _report_blocks(title, message):
        kind, text = block["type"], block.get("text", "")
        if kind == "heading":
            doc.add_paragraph(_plain_inline(text), style=f"Heading {min(block['level'], 3)}")
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet"); add_inline(p, text)
        elif kind == "numbered":
            p = doc.add_paragraph(style="List Number"); add_inline(p, re.sub(r"^\d+[.)]\s+", "", text))
        elif kind == "quote":
            paragraph = doc.add_paragraph(style="Quote"); add_inline(paragraph, text)
            paragraph.paragraph_format.left_indent = Inches(0.25)
            for run in paragraph.runs: run.font.color.rgb = RGBColor(71, 85, 105)
        elif kind == "code":
            try:
                from pygments import lex
                from pygments.lexers import get_lexer_by_name
                lexer = get_lexer_by_name(block.get("language") or "text", stripnl=False)
                tokens = list(lex(text, lexer))
            except Exception:
                tokens = [(None, text)]
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.12)
            paragraph.paragraph_format.right_indent = Inches(0.12)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(9)
            ppr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F1F5F9"); ppr.append(shd)
            for token, value in tokens:
                pieces = value.split("\n")
                for idx, piece in enumerate(pieces):
                    if piece:
                        run = paragraph.add_run(piece); run.font.name = "Consolas"; run.font.size = Pt(8.5)
                        color = getattr(token, "color", None) if token else None
                        if color: run.font.color.rgb = RGBColor.from_string(color)
                    if idx < len(pieces)-1: paragraph.add_run().add_break()
        elif kind == "table":
            headers, rows = block["headers"], block["rows"]
            table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"
            table.autofit = True
            header_row = table.rows[0]
            header_props = header_row._tr.get_or_add_trPr(); repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); header_props.append(repeat)
            for idx, value in enumerate(headers):
                cell = header_row.cells[idx]; shade(cell, "194A89")
                p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                run = p.add_run(_plain_inline(value)); run.bold = True; run.font.color.rgb = RGBColor(255,255,255); run.font.size = Pt(9)
            for row_idx, values in enumerate(rows):
                cells = table.add_row().cells
                for idx, value in enumerate(values):
                    if row_idx % 2 == 1: shade(cells[idx], "F1F5F9")
                    p = cells[idx].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                    run = p.add_run(_plain_inline(value)); run.font.size = Pt(9)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        elif kind == "math":
            image_data, image_width, image_height, _display = math_images[text]
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_together = True
            paragraph.add_run().add_picture(io.BytesIO(image_data), width=Inches(min(5.9, image_width / 192)))
        elif kind == "aligned":
            paragraph = doc.add_paragraph()
            paragraph.alignment = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
            }[block["alignment"]]
            add_inline(paragraph, text)
        elif kind == "diagram":
            diagram = _flowchart_image(text)
            if diagram:
                paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(io.BytesIO(diagram), width=Inches(6.0))
            else:
                paragraph = doc.add_paragraph(style="Quote"); paragraph.add_run(text).font.name = "Consolas"
        elif kind == "rule":
            paragraph = doc.add_paragraph(); ppr = paragraph._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom"); bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6"); bottom.set(qn("w:color"), "CBD5E1"); border.append(bottom); ppr.append(border)
        elif text:
            paragraph = doc.add_paragraph(); add_inline(paragraph, text)

    sources = _source_labels(message)
    if sources:
        doc.add_heading("Sources", level=1)
        for source in sources:
            p = doc.add_paragraph(style="List Bullet"); p.add_run(source)

    footer = section.footer.paragraphs[0]
    footer.style = "Caption"
    footer.paragraph_format.tab_stops.add_tab_stop(
        section.page_width - section.left_margin - section.right_margin,
        WD_TAB_ALIGNMENT.RIGHT,
    )
    brand_run = footer.add_run("PwaniNet  ·  Pwanimate")
    brand_run.font.color.rgb = RGBColor(100, 116, 139)
    footer.add_run("\t")
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _render_pdf(title, message):
    from playwright.sync_api import sync_playwright

    page_html, markdown, diagram_replacements = _render_markdown_html(title, message)
    with sync_playwright() as playwright:
        try:
            browser = playwright.chromium.launch(channel="chrome", headless=True, args=["--no-sandbox", "--disable-setuid-sandbox"])
        except Exception:
            try:
                browser = playwright.chromium.launch(headless=True, args=["--no-sandbox", "--disable-setuid-sandbox"])
            except Exception as exc:
                raise RuntimeError("PDF export requires Google Chrome or the Playwright Chromium browser.") from exc

        try:
            page = browser.new_page(viewport={"width": 1240, "height": 1754}, color_scheme="light")
            page.route("https://**/*", lambda route: route.abort())
            page.route("http://**/*", lambda route: route.abort())
            page.set_content(page_html, wait_until="load")
            for asset_key in ("marked", "dompurify", "katex", "prism", "prism_python", "prism_sql", "prism_bash", "prism_json", "prism_javascript", "prism_markup"):
                asset_path = _static_asset({
                    "marked": "vendor/marked/marked.min.js", "dompurify": "vendor/dompurify/purify.min.js",
                    "katex": "vendor/katex/katex.min.js", "prism": "vendor/prism/prism.min.js",
                    "prism_python": "vendor/prism/prism-python.min.js", "prism_sql": "vendor/prism/prism-sql.min.js",
                    "prism_bash": "vendor/prism/prism-bash.min.js", "prism_json": "vendor/prism/prism-json.min.js",
                    "prism_javascript": "vendor/prism/components/prism-javascript.min.js",
                    "prism_markup": "vendor/prism/components/prism-markup.min.js",
                }[asset_key])
                page.add_script_tag(content=asset_path.read_text(encoding="utf-8"))

            rendered = page.evaluate(r"""([title, markdown, diagramReplacements]) => {
                window.marked.use({ breaks: true, gfm: true });
                const placeholders = [];
                const storeMath = (code, displayMode) => {
                    const id = `PWANIMATE_MATH_PH_${placeholders.length}_XYZ`;
                    placeholders.push({ id, code, displayMode });
                    return id;
                };
                let processed = markdown || '';
                processed = processed.replace(/\$\$([\s\S]*?)\$\$/g, (_m, code) => `\n\n${storeMath(code, true)}\n\n`);
                processed = processed.replace(/\\\[([\s\S]*?)\\\]/g, (_m, code) => `\n\n${storeMath(code, true)}\n\n`);
                processed = processed.replace(/\\\(([\s\S]*?)\\\)/g, (_m, code) => storeMath(code, false));
                processed = processed.replace(/(?<!\\)\$([^$\n]+?)(?<!\\)\$/g, (_m, code) => code.trim() ? storeMath(code, false) : _m);
                let html = window.marked.parse(processed);
                html = window.DOMPurify.sanitize(html, {
                    ADD_TAGS: ['annotation','math','semantics','mrow','mi','mn','mo','msup','msub','mfrac','mover','munder','msqrt','mroot','mtd','mtr','mtable','mtext','mspace'],
                    ADD_ATTR: ['xmlns','display','displaystyle','mathvariant','columnalign','rowalign','linethickness']
                });
                for (const [token, safeHtml] of Object.entries(diagramReplacements || {})) html = html.split(token).join(safeHtml);
                for (const {id, code, displayMode} of placeholders) {
                    try {
                        const decoded = code.replace(/&amp;/g,'&').replace(/&lt;/g,'<').replace(/&gt;/g,'>').replace(/&quot;/g,'"').replace(/&#39;/g,"'");
                        const renderedMath = window.katex.renderToString(decoded, {displayMode, throwOnError:false, output:'htmlAndMathml'});
                        html = html.split(id).join(renderedMath);
                    } catch (error) {
                        const fallback = displayMode ? `<pre class="katex-fallback">${window.DOMPurify.sanitize(code)}</pre>` : `<code>${window.DOMPurify.sanitize(code)}</code>`;
                        html = html.split(id).join(fallback);
                    }
                }
                document.querySelector('.export-title').textContent = title;
                const container = document.getElementById('answer');
                container.innerHTML = html;
                container.querySelectorAll('table').forEach(table => {
                    const wrapper = document.createElement('div');
                    wrapper.className = 'pwanimate-table-scroll pwanimate-table-wrapper';
                    table.parentNode.insertBefore(wrapper, table);
                    wrapper.appendChild(table);
                    const columns = table.querySelector('tr')?.children.length || 0;
                    if (columns >= 5) wrapper.classList.add('wide-table-page');
                    if (columns >= 9) wrapper.classList.add('dense-wide-table');
                });
                container.querySelectorAll('pre').forEach(pre => {
                    if (pre.classList.contains('katex-fallback')) return;
                    const code = pre.querySelector('code');
                    let language = code?.className.match(/language-([a-zA-Z0-9_+#-]+)/)?.[1] || 'code';
                    const wrapper = document.createElement('div'); wrapper.className = 'pwanimate-code-block';
                    const header = document.createElement('div'); header.className = 'pwanimate-code-header';
                    const label = document.createElement('span'); label.className = 'pwanimate-code-lang'; label.textContent = language;
                    header.appendChild(label); pre.parentNode.insertBefore(wrapper, pre); wrapper.append(header, pre);
                    if (code && window.Prism) { try { window.Prism.highlightElement(code); } catch (_) {} }
                });
                return true;
            }""", [title, markdown, diagram_replacements])
            if not rendered:
                raise RuntimeError("The response could not be rendered for PDF export.")
            page.evaluate("async () => { await document.fonts.ready; await new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve))); }")
            page.emulate_media(media="print")
            footer = """<div style='width:100%;box-sizing:border-box;padding:0 19mm;font-family:Arial,sans-serif;font-size:8pt;color:#64748b;display:flex;justify-content:space-between;align-items:center;'><span>PwaniNet · Pwanimate</span><span>Page <span class='pageNumber'></span></span></div>"""
            return page.pdf(
                format="A4", prefer_css_page_size=True, print_background=True,
                display_header_footer=True, footer_template=footer,
                margin={"top": "18mm", "right": "19mm", "bottom": "23mm", "left": "19mm"},
            )
        finally:
            browser.close()


def _document_title(message):
    generic_titles = {
        "answer", "overview", "summary", "introduction", "key points", "conclusion",
        "response", "study notes", "report", "untitled",
    }
    headings = [block for block in _content_blocks(message.content) if block["type"] == "heading" and block["level"] <= 2]
    for block in headings:
        candidate = re.sub(r"\s+", " ", _plain_inline(block["text"])).strip(" #\t\r\n")
        if candidate and candidate.casefold().strip(" .:;-") not in generic_titles:
            title = candidate
            break
    else:
        prompt = message.conversation.messages.filter(role="user").order_by("-created_at").first()
        title = _clean_inline(prompt.content) if prompt else "Study Notes"
        title = re.split(r"(?<=[.!?])\s+", title, maxsplit=1)[0]
        title = re.sub(r"^(?:please\s+)?(?:can you|could you|would you|please)\s+", "", title, flags=re.IGNORECASE)
        title = re.sub(r"\s+", " ", title).strip(" \t\r\n?!.:")
        if len(title) > 100:
            title = title[:100].rsplit(" ", 1)[0].rstrip(" ,;:-")
        if title:
            title = title[0].upper() + title[1:]
        else:
            title = "Study Notes"
    return title[:180] or "Study Notes"


@transaction.atomic
def create_generated_resource(message, user, file_format):
    if file_format not in {"pdf", "docx"}:
        raise ValueError("Choose PDF or Word format.")
    content = (message.content or "").strip()
    if not content:
        raise ValueError("There is no response content to save.")
    if len(content) > MAX_RESOURCE_CHARACTERS:
        raise ValueError("This response is too long to export as one document.")

    existing = Document.objects.filter(
        uploaded_by=user,
        is_ai_generated=True,
        generated_from_message_id=message.id,
        generated_format=file_format,
    ).select_related("category").first()
    if existing:
        existing_file = existing.latest_version.files.first()
        if existing.description == GENERATED_RESOURCE_DESCRIPTION and existing_file:
            return existing, existing_file, False

    title = _document_title(message)
    category, _ = Category.objects.get_or_create(
        code="ai_generated",
        defaults={"name": "AI Generated Resource", "icon": "bi-stars"},
    )
    base_slug = slugify(title)[:500] or "pwanimate-resource"
    slug = base_slug
    suffix = 2
    while Document.objects.filter(slug=slug).exists():
        slug = f"{base_slug[:480]}-{suffix}"
        suffix += 1

    document = Document.objects.create(
        title=title,
        slug=slug,
        description=GENERATED_RESOURCE_DESCRIPTION,
        category=category,
        uploaded_by=user,
        visibility="private",
        status="processing",
        is_available=True,
        is_ai_generated=True,
        generated_from_message_id=message.id,
        generated_format=file_format,
    )
    version = DocumentVersion.objects.create(
        document=document,
        version_number=1,
        change_notes="Generated by Pwanimate",
        created_by=user,
        is_latest=True,
    )

    data = _render_pdf(title, message) if file_format == "pdf" else _render_docx(title, message)
    extension = "pdf" if file_format == "pdf" else "docx"
    mime_type = "application/pdf" if extension == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    filename = f"{slug[:160]}.{extension}"
    private_storage = get_private_resource_storage()
    if existing:
        resource_file = existing.latest_version.files.first()
        if not resource_file:
            raise RuntimeError("The existing generated resource has no stored file to refresh.")
        old_storage_path = resource_file.file.name
        checksum = hashlib.sha256(data).hexdigest()
        versioned_filename = f"{slug[:140]}-{checksum[:12]}.{extension}"
        resource_file.file.storage = private_storage
        resource_file.file.save(versioned_filename, ContentFile(data), save=False)
        resource_file.storage_path = resource_file.file.name
        resource_file.original_filename = filename
        resource_file.mime_type = mime_type
        resource_file.extension = extension
        resource_file.size_bytes = len(data)
        resource_file.checksum = checksum
        resource_file.storage_provider = "r2" if private_storage.uses_s3 else "local"
        resource_file.processing_status = "pending"
        resource_file.storage_private = True
        resource_file.save(update_fields=[
            "file", "storage_path", "original_filename", "mime_type", "extension",
            "size_bytes", "checksum", "storage_provider", "processing_status", "storage_private",
        ])
        existing.description = GENERATED_RESOURCE_DESCRIPTION
        existing.save(update_fields=["description"])

        def cleanup_old_file():
            if old_storage_path and old_storage_path != resource_file.file.name:
                private_storage.delete(old_storage_path)

        def queue_refresh():
            from documents.tasks.processing import process_document
            process_document.delay(existing.id)

        transaction.on_commit(cleanup_old_file)
        transaction.on_commit(queue_refresh)
        return existing, resource_file, True

    resource_file = DocumentFile.objects.create(
        document_version=version,
        original_filename=filename,
        storage_path="",
        mime_type=mime_type,
        extension=extension,
        size_bytes=len(data),
        checksum=hashlib.sha256(data).hexdigest(),
        storage_provider="r2" if private_storage.uses_s3 else "local",
        processing_status="pending",
        uploaded_by=user,
        storage_private=True,
    )
    resource_file.file.storage = private_storage
    resource_file.file.save(filename, ContentFile(data), save=False)
    resource_file.storage_path = resource_file.file.name
    resource_file.save(update_fields=["file", "storage_path"])

    def queue_processing():
        from documents.tasks.processing import process_document
        process_document.delay(document.id)

    transaction.on_commit(queue_processing)
    return document, resource_file, True
