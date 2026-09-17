"""
Slide-aware PPTX structural extractor using python-pptx.
Preserves slide numbers, titles, body text, tables, and speaker notes.
"""

import os
import logging
from typing import Optional
from .base import BaseDocumentExtractor
from ..types import ExtractedDocument, ExtractedElement, ElementType, LocationMetadata
from ..exceptions import CorruptFileError, EmptyDocumentError, ExtractionError

logger = logging.getLogger(__name__)


class PPTXExtractor(BaseDocumentExtractor):
    """
    Extracts presentation slides as distinct structural units, capturing
    slide titles, body contents, tables, and presenter notes per slide.
    """

    def extract(self, file_path: str, filename: Optional[str] = None, **kwargs) -> ExtractedDocument:
        if not file_path or not os.path.isfile(file_path):
            raise CorruptFileError(f"PPTX file not found: {file_path}", source_filename=filename or file_path, format="pptx")

        display_name = filename or os.path.basename(file_path)

        try:
            import pptx
        except ImportError as e:
            raise ExtractionError("python-pptx library is not installed", source_filename=display_name, format="pptx") from e

        try:
            prs = pptx.Presentation(file_path)
        except Exception as e:
            raise CorruptFileError(f"Cannot open PPTX presentation: {e}", source_filename=display_name, format="pptx") from e

        total_slides = len(prs.slides)
        if total_slides == 0:
            raise EmptyDocumentError(f"Presentation has zero slides: {display_name}", source_filename=display_name, format="pptx")

        elements = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1

            # 1. Identify slide title if present
            slide_title = ""
            title_shape = None
            try:
                if slide.shapes.title and slide.shapes.title.text:
                    slide_title = slide.shapes.title.text.replace('\x00', '').strip()
                    title_shape = slide.shapes.title
            except Exception:
                title_shape = None

            if slide_title:
                elements.append(
                    ExtractedElement(
                        element_type=ElementType.HEADING,
                        text=slide_title,
                        location=LocationMetadata(
                            slide_number=slide_num,
                            section_heading=slide_title,
                        ),
                        metadata={"is_slide_title": True},
                    )
                )

            # 2. Extract shape text frames and tables in slide order
            for shape_idx, shape in enumerate(slide.shapes):
                if shape == title_shape:
                    continue  # Already captured as slide title

                if shape.has_text_frame:
                    paragraphs = []
                    for p in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in p.runs).replace('\x00', '').strip()
                        if text:
                            paragraphs.append(text)

                    if paragraphs:
                        combined_para = "\n".join(paragraphs)
                        elements.append(
                            ExtractedElement(
                                element_type=ElementType.PARAGRAPH,
                                text=combined_para,
                                location=LocationMetadata(
                                    slide_number=slide_num,
                                    section_heading=slide_title or None,
                                    block_index=shape_idx,
                                ),
                                metadata={"shape_type": "text_frame"},
                            )
                        )

                elif shape.has_table:
                    table_rows = []
                    for row in shape.table.rows:
                        cells = [c.text.replace('\x00', '').strip() for c in row.cells]
                        if any(cells):
                            table_rows.append(" | ".join(cells))

                    if table_rows:
                        elements.append(
                            ExtractedElement(
                                element_type=ElementType.TABLE,
                                text="\n".join(table_rows),
                                location=LocationMetadata(
                                    slide_number=slide_num,
                                    section_heading=slide_title or None,
                                    block_index=shape_idx,
                                ),
                                metadata={"shape_type": "table", "row_count": len(table_rows)},
                            )
                        )

            # 3. Extract speaker notes if available
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes_text = slide.notes_slide.notes_text_frame.text.replace('\x00', '').strip()
                    if notes_text:
                        elements.append(
                            ExtractedElement(
                                element_type=ElementType.NOTES,
                                text=notes_text,
                                location=LocationMetadata(
                                    slide_number=slide_num,
                                    section_heading=slide_title or None,
                                ),
                                metadata={"is_speaker_notes": True},
                            )
                        )
            except Exception as notes_err:
                logger.debug(f"Could not read speaker notes for slide {slide_num}: {notes_err}")

        if not elements:
            raise EmptyDocumentError(f"No extractable text found in PPTX: {display_name}", source_filename=display_name, format="pptx")

        return ExtractedDocument(
            source_filename=display_name,
            format="pptx",
            elements=elements,
            metadata={
                "total_slides": total_slides,
                "element_count": len(elements),
            }
        )
